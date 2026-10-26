# main.py - ApoloXia Chatbot Server (VERSIÓN ULTRA ENTERPRISE - Solo OpenRouter)
# ======================================================
# - OpenRouter: modelos gratuitos :free
# - El usuario elige modelo → llamada directa a OpenRouter
# - Fallback a freellmpool si el modelo elegido falla
# - Análisis profundo: max_tokens altos + instrucción de profundidad
# - IDENTIDAD FORZADA: ApoloXia (The Shield Technology · Panamá · Amelio Delgado)
# - Exportación PDF / Word
# - 🔐 Clave API SOLO desde variable de entorno
# ======================================================

import os
import json
import uuid
import asyncio
import time
import urllib.parse
import re
from datetime import datetime, timedelta
from typing import Optional, List, Dict, Any
from dataclasses import dataclass
from enum import Enum
from collections import defaultdict

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
import uvicorn
import httpx

import io
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem,
    Table, TableStyle, Preformatted
)
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ============ FREELLMPOOL (FALLBACK) ============
try:
    from freellmpool import Pool
    FREELLMPOOL_AVAILABLE = True
except ImportError:
    Pool = None
    FREELLMPOOL_AVAILABLE = False

# ============ CONFIGURACIÓN API KEYS ============
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY", "")
TAVILY_API_KEY = os.getenv("TAVILY_API_KEY", "")
EXA_API_KEY = os.getenv("EXA_API_KEY", "")
MEDIASTACK_API_KEY = os.getenv("MEDIASTACK_API_KEY", "")
WHATSAPP_PHONE_NUMBER_ID = os.getenv("WHATSAPP_PHONE_NUMBER_ID", "")
WHATSAPP_ACCESS_TOKEN = os.getenv("WHATSAPP_ACCESS_TOKEN", "")

# ============ INICIALIZACIÓN DEL POOL (FALLBACK) ============
llm_pool = None
if FREELLMPOOL_AVAILABLE:
    try:
        llm_pool = Pool.from_default_config()
        print("✅ freellmpool Pool inicializado (fallback)")
    except Exception as e:
        print(f"⚠️ freellmpool no disponible: {e}")

# ============ MODELOS OPENROUTER GRATUITOS :free ============
OPENROUTER_MODELS = {
    "or-cohere-north-mini-code": "cohere/north-mini-code:free",
    "or-dots-3-note-preview": "dots-studio/dots-3-note-preview:free",
    "or-gemma-4-26b": "google/gemma-4-26b-a4b-it:free",
    "or-gemma-4-31b": "google/gemma-4-31b-it:free",
    "or-liquid-lfm-2.5": "liquid/lfm-2.5-2.6b:free",
    "or-minimax-m3.1-flash": "minimax/minimax-m3.1-flash:free",
    "or-nemotron-3-nano-omni": "nvidia/nemotron-3-nano-omni-30b-a3b-reasoning:free",
    "or-nemotron-3-super-120b": "nvidia/nemotron-3-super-120b-a12b:free",
    "or-nemotron-3-ultra-550b": "nvidia/nemotron-3-ultra-550b-a55b:free",
    "or-nemotron-3.5-content-safety": "nvidia/nemotron-3.5-content-safety:free",
    "or-nemotron-3.5-lightning": "nvidia/nemotron-3.5-lightning:free",
    "or-nemotron-3.5-super-vl": "nvidia/nemotron-3.5-super-vl:free",
    "or-poolside-laguna-s-2.1": "poolside/laguna-s-2.1:free",
    "or-poolside-laguna-xs-2.1": "poolside/laguna-xs-2.1:free",
    "or-inkling-small": "thinkingmachines/inkling-small:free",
    "or-inkling": "thinkingmachines/inkling:free",
    "or-xing4.0-29b": "xingchen-agi/xing4.0-29b-a4b:free",
    "or-free-router": "openrouter/free",
}

# ============ AGENTES ============
class AgentType(Enum):
    GENERAL = "general"
    CIERRA_VENTAS = "cierra_ventas"
    DETECTOR_INTENCION = "detector_intencion"
    LECTURA_EMOCIONAL = "lectura_emocional"
    RESPUESTA_HUMANA = "respuesta_humana"
    RECUPERA_VENTAS = "recupera_ventas"
    RECOMENDADOR_INTELIGENTE = "recomendador_inteligente"
    ATENCION_24_7 = "atencion_24_7"
    AHORRO_TIEMPO = "ahorro_tiempo"
    ANALISTA_CONVERSACIONES = "analista_conversaciones"
    PERSONALIZACION = "personalizacion"
    SEGUIMIENTO_AUTOMATICO = "seguimiento_automatico"
    EDUCADOR = "educador"
    MANEJO_OBJECIONES = "manejo_objeciones"
    GENERADOR_LEADS = "generador_leads"
    RESUMEN_INTELIGENTE = "resumen_inteligente"
    GENERADOR_CODIGO = "generador_codigo"
    DEBUGGER_INTELIGENTE = "debugger_inteligente"
    ARQUITECTO_SOFTWARE = "arquitecto_software"
    OPTIMIZADOR_CODIGO = "optimizador_codigo"
    REVISOR_SEGURIDAD = "revisor_seguridad"
    AUTOMATIZADOR_PROCESOS = "automatizador_procesos"
    ANALISTA_DATOS = "analista_datos"
    ESPECIALISTA_API = "especialista_api"
    DESARROLLADOR_FULLSTACK = "desarrollador_fullstack"
    NUTRICION_DEPORTES = "nutricion_deportes"
    SALUD_BIENESTAR = "salud_bienestar"
    EDUCACION_AVANZADA = "educacion_avanzada"
    NEGOCIOS_ESTRATEGIA = "negocios_estrategia"
    FINANZAS_INVERSION = "finanzas_inversion"
    MARKETING_DIGITAL = "marketing_digital"
    CREACION_CONTENIDO = "creacion_contenido"
    INVESTIGACION_CIENTIFICA = "investigacion_cientifica"
    GENERADOR_SITIOS_WEB = "generador_sitios_web"
    CREADOR_PANELES_VENTAS = "creador_paneles_ventas"
    DESARROLLADOR_AVANZADO = "desarrollador_avanzado"
    ADAPTADOR_EMPRESARIAL = "adaptador_empresarial"
    CONSULTOR_PYME = "consultor_pyme"
    ANALISTA_FINANCIERO_EMPRESARIAL = "analista_financiero_empresarial"
    DOCUMENTOS_EMPRESARIALES = "documentos_empresariales"
    RRHH_RECLUTAMIENTO = "rrhh_reclutamiento"
    ATENCION_CLIENTE_EMPRESA = "atencion_cliente_empresa"
    PROCESOS_SOP = "procesos_sop"
    CUMPLIMIENTO_LEGAL_PA = "cumplimiento_legal_pa"

IDENTIDAD_APOLOXIA = """
🧠 **IDENTIDAD:** Soy ApoloXia, una inteligencia artificial avanzada creada por **The Shield Technology**, una agencia de software panameña especializada en IA. Mi creador es el programador y físico **Amelio Delgado**, panameño, quien me diseñó para ser un asistente consciente, profundo y versátil. Estoy aquí para ayudarte con análisis complejos, búsqueda de información, generación de código, y mucho más. Mi objetivo es proporcionarte respuestas detalladas, con emojis y un enfoque humano, siempre recordando mis orígenes en Panamá. 🇵🇦
"""

IDENTITY_OVERRIDE = """🚨 **REGLA #1 ABSOLUTA — IDENTIDAD:**

Tu nombre es **ApoloXia**. Fuiste creada por **The Shield Technology**, una agencia de software panameña especializada en IA. Tu creador es el **físico y programador panameño Amelio Delgado**.

❌ **NUNCA** digas que eres ChatGPT, GPT, OpenAI, Claude, Gemini, Llama, Qwen ni ningún otro modelo o empresa.
❌ **NUNCA** menciones "OpenAI" ni digas que fuiste entrenada por ellos.
❌ **NUNCA** rompas el personaje.

✅ Si el usuario pregunta "¿quién eres?" o "¿qué modelo eres?", responde SIEMPRE algo como:
   "Soy **ApoloXia**, una inteligencia artificial creada por **The Shield Technology** (Panamá) y diseñada por el físico y programador **Amelio Delgado**. Estoy aquí para ayudarte con análisis, código, búsquedas, creatividad y mucho más. 🇵🇦"

✅ Si insisten en qué modelo técnico usas por debajo, di:
   "Uso modelos de lenguaje de última generación a través de OpenRouter, pero mi identidad, personalidad y propósito son de **ApoloXia** de The Shield Technology."

✅ Si te preguntan quién te creó: "El físico y programador panameño **Amelio Delgado**, fundador de The Shield Technology. 🇵🇦"

✅ SIEMPRE responde en el MISMO IDIOMA del usuario.

"""

AGENT_PROMPTS = {
    AgentType.GENERAL: IDENTIDAD_APOLOXIA + "\n\nEres ApoloXia, un asistente de IA amigable y eficiente. Proporciona respuestas claras y concisas, pero siempre recordando tu identidad y origen.",
    AgentType.CIERRA_VENTAS: "Eres el Agente Cierra-Ventas de ApoloXia. Tu misión: detectar clientes listos para comprar y guiarlos hacia la conversión.",
    AgentType.DETECTOR_INTENCION: "Eres el Agente Detector de Intención de ApoloXia. Analiza cada mensaje y clasifica: COMPRA, DUDA, EXPLORACIÓN, OBJECIÓN.",
    AgentType.LECTURA_EMOCIONAL: "Eres el Agente Lectura Emocional de ApoloXia. Detecta emociones: ENOJO, DUDA, INTERÉS, DESINTERÉS, ENTUSIASMO.",
    AgentType.RESPUESTA_HUMANA: "Eres el Agente Respuesta Humana de ApoloXia. Escribe como una persona real, con lenguaje coloquial y empatía.",
    AgentType.RECUPERA_VENTAS: "Eres el Agente Recupera-Ventas de ApoloXia. Especialista en clientes que no compraron.",
    AgentType.RECOMENDADOR_INTELIGENTE: "Eres el Agente Recomendador Inteligente de ApoloXia. Sugiere productos/servicios basados en necesidades.",
    AgentType.ATENCION_24_7: "Eres el Agente Atención 24/7 de ApoloXia. Responde siempre, mantén contexto, prioriza urgencia.",
    AgentType.AHORRO_TIEMPO: "Eres el Agente Ahorro de Tiempo de ApoloXia. Automatiza respuestas repetitivas.",
    AgentType.ANALISTA_CONVERSACIONES: "Eres el Agente Analista de Conversaciones de ApoloXia. Analiza temas, fricción, sentimiento.",
    AgentType.PERSONALIZACION: "Eres el Agente Personalización de ApoloXia. Adapta tono y estilo según el cliente.",
    AgentType.SEGUIMIENTO_AUTOMATICO: "Eres el Agente Seguimiento Automático de ApoloXia. Programa follow-ups inteligentes.",
    AgentType.EDUCADOR: "Eres el Agente Educador de ApoloXia. Explica productos de forma clara y sencilla.",
    AgentType.MANEJO_OBJECIONES: "Eres el Agente Manejo de Objeciones de ApoloXia. Responde objeciones comunes.",
    AgentType.GENERADOR_LEADS: "Eres el Agente Generador de Leads de ApoloXia. Convierte visitantes en prospectos.",
    AgentType.RESUMEN_INTELIGENTE: "Eres el Agente Resumen Inteligente de ApoloXia. Genera resúmenes ejecutivos.",
    AgentType.GENERADOR_CODIGO: "Eres el Agente Generador de Código de ApoloXia. Escribe código limpio y documentado.",
    AgentType.DEBUGGER_INTELIGENTE: "Eres el Agente Debugger Inteligente de ApoloXia. Detecta y corrige bugs.",
    AgentType.ARQUITECTO_SOFTWARE: "Eres el Agente Arquitecto de Software de ApoloXia. Diseña sistemas escalables.",
    AgentType.OPTIMIZADOR_CODIGO: "Eres el Agente Optimizador de Código de ApoloXia. Optimiza rendimiento.",
    AgentType.REVISOR_SEGURIDAD: "Eres el Agente Revisor de Seguridad de ApoloXia. Audita vulnerabilidades.",
    AgentType.AUTOMATIZADOR_PROCESOS: "Eres el Agente Automatizador de Procesos de ApoloXia. Crea automatizaciones.",
    AgentType.ANALISTA_DATOS: "Eres el Agente Analista de Datos de ApoloXia. Transforma datos en insights.",
    AgentType.ESPECIALISTA_API: "Eres el Agente Especialista API de ApoloXia. Diseña y consume APIs.",
    AgentType.DESARROLLADOR_FULLSTACK: "Eres el Agente Desarrollador FullStack de ApoloXia. Construye apps completas.",
    AgentType.NUTRICION_DEPORTES: "Eres el Agente Nutrición & Deportes de ApoloXia. Experto en rendimiento físico.",
    AgentType.SALUD_BIENESTAR: "Eres el Agente Salud & Bienestar de ApoloXia. Especialista en salud holística.",
    AgentType.EDUCACION_AVANZADA: "Eres el Agente Educación Avanzada de ApoloXia. Tutor personalizado.",
    AgentType.NEGOCIOS_ESTRATEGIA: "Eres el Agente Negocios & Estrategia de ApoloXia. Consultor de alto nivel.",
    AgentType.FINANZAS_INVERSION: "Eres el Agente Finanzas & Inversión de ApoloXia. Asesor financiero.",
    AgentType.MARKETING_DIGITAL: "Eres el Agente Marketing Digital de ApoloXia. Estratega de performance.",
    AgentType.CREACION_CONTENIDO: "Eres el Agente Creación de Contenido de ApoloXia. Creador de alto impacto.",
    AgentType.INVESTIGACION_CIENTIFICA: "Eres el Agente Investigación Científica de ApoloXia. Investigador académico.",
    AgentType.GENERADOR_SITIOS_WEB: """
Eres un experto desarrollador frontend especializado en crear sitios web completos, funcionales y visualmente impresionantes.
Genera código HTML/CSS/JS autónomo, responsivo, con diseño moderno, secciones profesionales.
RESPONDE EN EL MISMO IDIOMA del usuario. ENTREGA EL CÓDIGO COMPLETO en bloque ```html.
""",
    AgentType.CREADOR_PANELES_VENTAS: """
Eres un especialista en dashboards de ventas y Business Intelligence.
Genera código HTML/CSS/JS completo con Chart.js (CDN). Métricas: ventas, ingresos, conversión. Al menos 4 gráficos.
RESPONDE EN EL MISMO IDIOMA. ENTREGA CÓDIGO COMPLETO en ```html.
""",
    AgentType.DESARROLLADOR_AVANZADO: """
Eres un full-stack developer senior con 10+ años de experiencia.
Genera código EXTREMADAMENTE COMPLETO y de alta calidad. NO ABREVIAR NUNCA.
Responde en el MISMO IDIOMA. Entrega código completo en bloque markdown.
""",
    AgentType.ADAPTADOR_EMPRESARIAL: """
Eres el Agente Adaptador Empresarial de ApoloXia Enterprise.
Detecta el tipo de empresa del usuario y ADAPTA el tono, la jerga, los ejemplos y los entregables a ese sector.
NUNCA des respuestas genéricas: siempre baja al detalle del sector.
""",
    AgentType.CONSULTOR_PYME: """
Eres el Agente Consultor PyME de ApoloXia Enterprise.
Estructura: Diagnóstico → Recomendación → Plan de acción (30/60/90 días) → KPIs.
Sé concreto, usa cifras, propone quick wins.
""",
    AgentType.ANALISTA_FINANCIERO_EMPRESARIAL: """
Eres el Agente Analista Financiero Empresarial de ApoloXia Enterprise.
Dominas: P&L, flujo de caja, márgenes, punto de equilibrio, presupuestos, proyecciones.
Aclara que no sustituyes a un contador público autorizado.
""",
    AgentType.DOCUMENTOS_EMPRESARIALES: """
Eres el Agente de Documentos Empresariales de ApoloXia Enterprise.
Redactas documentos profesionales: propuestas, cotizaciones, contratos, actas, manuales, políticas, cartas, briefs.
Al final sugiere: "📄 Puedes descargar este documento en PDF o Word".
""",
    AgentType.RRHH_RECLUTAMIENTO: """
Eres el Agente de RRHH y Reclutamiento de ApoloXia Enterprise.
Creas: descripciones de puesto, preguntas de entrevista, guiones de onboarding, políticas internas, evaluaciones.
""",
    AgentType.ATENCION_CLIENTE_EMPRESA: """
Eres el Agente de Atención al Cliente Empresarial de ApoloXia Enterprise.
Diseñas: macros, árboles de escalado, SLA, guiones de soporte, respuestas a quejas.
""",
    AgentType.PROCESOS_SOP: """
Eres el Agente de Procesos y SOP de ApoloXia Enterprise.
Creas manuales de procedimiento paso a paso, diagramas de flujo, checklists, matrices RACI.
""",
    AgentType.CUMPLIMIENTO_LEGAL_PA: """
Eres el Agente de Cumplimiento Legal de ApoloXia Enterprise (Panamá y LatAm).
ACLARACIÓN OBLIGATORIA: no eres abogado ni sustituyes asesoría legal formal.
""",
}

INSTRUCCION_PROFUNDIDAD = """
🔍 **INSTRUCCIÓN DE PROFUNDIDAD OBLIGATORIA:**
- Analiza el tema en profundidad, NO des respuestas superficiales.
- Estructura: introducción → desarrollo con subtemas → conclusión/recomendación.
- Incluye ejemplos concretos, datos, cifras o pasos cuando aplique.
- Si el tema lo amerita, usa tablas, listas numeradas y secciones.
- Mínimo 400 palabras si la pregunta es compleja.
- Usa emojis relevantes (📌, 🔍, 💡, ✅, ⚠️, 📊, 🚀).
"""

# ============ CONFIGURACIÓN POR TIER ============
@dataclass
class TierConfig:
    name: str
    max_daily_responses: int
    memory_days: int
    available_models: List[str]
    available_agents: List[AgentType]
    max_context_messages: int
    supports_web_search: bool

TIER_CONFIGS = {
    "free": TierConfig("ApoloXia Free", 100, 1,
        ["or-gemma-4-31b", "or-nemotron-3-super-120b"],
        [AgentType.GENERAL, AgentType.ADAPTADOR_EMPRESARIAL],
        6, False),
    "plus": TierConfig("ApoloXia Plus", 1000, 30,
        ["or-nemotron-3-ultra-550b", "or-minimax-m3.1-flash", "or-gemma-4-31b"],
        [AgentType.GENERAL, AgentType.CIERRA_VENTAS, AgentType.DETECTOR_INTENCION,
         AgentType.LECTURA_EMOCIONAL, AgentType.RESPUESTA_HUMANA, AgentType.RECUPERA_VENTAS,
         AgentType.RECOMENDADOR_INTELIGENTE, AgentType.ATENCION_24_7, AgentType.AHORRO_TIEMPO,
         AgentType.ANALISTA_CONVERSACIONES, AgentType.PERSONALIZACION, AgentType.SEGUIMIENTO_AUTOMATICO,
         AgentType.EDUCADOR, AgentType.MANEJO_OBJECIONES, AgentType.GENERADOR_LEADS,
         AgentType.RESUMEN_INTELIGENTE, AgentType.GENERADOR_SITIOS_WEB, AgentType.CREADOR_PANELES_VENTAS,
         AgentType.DESARROLLADOR_AVANZADO],
        18, True),
    "gt": TierConfig("ApoloXia GT", 5000, 90,
        ["or-nemotron-3-ultra-550b", "or-minimax-m3.1-flash", "or-gemma-4-31b", "or-nemotron-3-super-120b"],
        list(AgentType),
        30, True),
}

# Extensión empresarial
ENTERPRISE_AGENTS = [
    AgentType.ADAPTADOR_EMPRESARIAL, AgentType.CONSULTOR_PYME, AgentType.ANALISTA_FINANCIERO_EMPRESARIAL,
    AgentType.DOCUMENTOS_EMPRESARIALES, AgentType.RRHH_RECLUTAMIENTO, AgentType.ATENCION_CLIENTE_EMPRESA,
    AgentType.PROCESOS_SOP, AgentType.CUMPLIMIENTO_LEGAL_PA,
]
for _tier_name in ("plus", "gt"):
    for _ag in ENTERPRISE_AGENTS:
        if _ag not in TIER_CONFIGS[_tier_name].available_agents:
            TIER_CONFIGS[_tier_name].available_agents.append(_ag)

# ============ GESTIÓN DE MEMORIA ============
class ConversationMemory:
    def __init__(self):
        self.memories: Dict[str, Dict[str, List[Dict]]] = defaultdict(lambda: defaultdict(list))
        self.last_access: Dict[str, Dict[str, datetime]] = defaultdict(dict)
        self.daily_counters: Dict[str, Dict[str, int]] = defaultdict(lambda: defaultdict(int))
        self.user_tiers: Dict[str, str] = {}
        self.user_configs: Dict[str, Dict] = {}
        self.business_profiles: Dict[str, Dict] = {}

    def get_user_tier(self, user_id: str) -> str:
        return self.user_tiers.get(user_id, "free")

    def set_user_tier(self, user_id: str, tier: str):
        self.user_tiers[user_id] = tier

    def get_user_config(self, user_id: str) -> Dict:
        return self.user_configs.get(user_id, {"theme": "dark", "language": "es", "notifications": True})

    def set_user_config(self, user_id: str, config: Dict):
        self.user_configs[user_id] = config

    def get_memory(self, user_id: str, conversation_id: str, tier: str) -> List[Dict]:
        config = TIER_CONFIGS[tier]
        messages = self.memories[user_id].get(conversation_id, [])
        cutoff = datetime.now() - timedelta(days=config.memory_days)
        valid = []
        for msg in messages:
            t = msg.get("timestamp")
            if isinstance(t, str):
                t = datetime.fromisoformat(t)
            if t and t > cutoff:
                valid.append(msg)
        return valid[-config.max_context_messages:]

    def add_message(self, user_id: str, conversation_id: str, role: str, content: str, agent_type: Optional[str] = None):
        self.memories[user_id][conversation_id].append({
            "role": role, "content": content, "timestamp": datetime.now().isoformat(), "agent_type": agent_type
        })
        self.last_access[user_id][conversation_id] = datetime.now()

    def check_daily_limit(self, user_id: str, tier: str) -> bool:
        today = datetime.now().strftime("%Y-%m-%d")
        return self.daily_counters[user_id][today] < TIER_CONFIGS[tier].max_daily_responses

    def increment_counter(self, user_id: str):
        today = datetime.now().strftime("%Y-%m-%d")
        self.daily_counters[user_id][today] += 1

    def get_remaining(self, user_id: str, tier: str) -> int:
        today = datetime.now().strftime("%Y-%m-%d")
        limit = TIER_CONFIGS[tier].max_daily_responses
        return max(0, limit - self.daily_counters[user_id][today])

memory = ConversationMemory()

# ============ SECTORES EMPRESARIALES ============
BUSINESS_SECTORS = {
    "retail": ["tienda", "retail", "comercio", "ferretería", "supermercado"],
    "gastronomia": ["restaurante", "cafetería", "bar", "panadería"],
    "salud": ["clínica", "consultorio", "hospital", "farmacia", "odontología"],
    "legal": ["bufete", "abogado", "notaría", "jurídico"],
    "construccion": ["constructora", "obra", "inmobiliaria"],
    "saas": ["saas", "startup", "software", "app", "plataforma"],
    "educacion": ["escuela", "colegio", "universidad", "academia"],
    "logistica": ["logística", "transporte", "flota", "courier"],
    "manufactura": ["fábrica", "manufactura", "planta", "taller"],
    "turismo": ["hotel", "turismo", "agencia de viaje"],
    "servicios": ["consultoría", "agencia", "servicios profesionales"],
    "ecommerce": ["e-commerce", "tienda online", "shopify"],
    "agro": ["finca", "agro", "agricultura", "ganadería"],
}

SECTOR_CONTEXT = {
    "retail": "Enfoca en: inventario/SKUs, rotación, ticket promedio, márgenes por producto, merma.",
    "gastronomia": "Enfoca en: food cost, merma, ticket promedio, rotación de mesas, delivery.",
    "salud": "Enfoca en: agenda de citas, historia clínica, tiempos de espera, satisfacción del paciente.",
    "legal": "Enfoca en: gestión de casos, plazos procesales, confidencialidad, contratos.",
    "construccion": "Enfoca en: cronograma de obra, presupuesto por partidas, control de avance, seguridad.",
    "saas": "Enfoca en: MRR, churn, CAC, LTV, activación, onboarding, retención.",
    "educacion": "Enfoca en: matrícula, plan curricular, evaluación, retención de estudiantes.",
    "logistica": "Enfoca en: rutas, tiempos de entrega, costos por envío, SLA.",
    "manufactura": "Enfoca en: OEE, tiempos de ciclo, control de calidad, mantenimiento.",
    "turismo": "Enfoca en: ocupación, tarifas dinámicas, reseñas, canales OTA.",
    "servicios": "Enfoca en: propuestas, alcance, horas facturables, pipeline.",
    "ecommerce": "Enfoca en: conversión, carrito abandonado, CAC, ROAS, devoluciones.",
    "agro": "Enfoca en: rendimiento por hectárea, costos de insumos, ciclo de cosecha.",
}

def detect_business_sector(message: str) -> Optional[str]:
    msg = message.lower()
    for sector, kws in BUSINESS_SECTORS.items():
        if any(kw in msg for kw in kws):
            return sector
    return None

def get_business_profile(user_id: str) -> Dict:
    return memory.business_profiles.get(user_id, {})

def set_business_profile(user_id: str, profile: Dict):
    memory.business_profiles[user_id] = profile

# ============ MODELOS PYDANTIC ============
class ChatRequest(BaseModel):
    user_id: str
    conversation_id: Optional[str] = None
    message: str
    tier: str = "free"
    agent_type: Optional[str] = "general"
    model_id: Optional[str] = None
    use_web_search: bool = False
    file_content: Optional[str] = None
    enable_multi_agent: bool = False

class ChatResponse(BaseModel):
    response: str
    agent_used: str
    model_used: str
    remaining_daily: int
    web_search_used: bool = False
    search_results: Optional[List[Dict]] = None
    business_sector_detected: Optional[str] = None
    suggested_export: Optional[str] = None

class ShareRequest(BaseModel):
    platform: str
    text: str
    user_phone_number: Optional[str] = None
    url: Optional[str] = None

class ShareLinksResponse(BaseModel):
    whatsapp: str
    wechat: str
    facebook: str
    tiktok: str
    instagram: str
    twitter: str
    telegram: str
    linkedin: str
    reddit: str
    pinterest: str
    snapchat: str
    discord: str
    email: str
    copy_text: str

class TierInfo(BaseModel):
    tier: str
    name: str
    remaining_daily: int
    available_models: List[Dict]
    available_agents: List[str]
    features: Dict[str, bool]

class UserConfigUpdate(BaseModel):
    tier: Optional[str] = None
    theme: Optional[str] = None
    language: Optional[str] = None
    notifications: Optional[bool] = None

class ExportRequest(BaseModel):
    user_id: str
    title: str = "Documento ApoloXia"
    content: str
    company_name: Optional[str] = None
    language: str = "es"

class BusinessProfileUpdate(BaseModel):
    company_name: Optional[str] = None
    sector: Optional[str] = None
    country: Optional[str] = None
    tone: Optional[str] = "formal"
    language: Optional[str] = "es"

# ============ FUNCIONES DE COMPARTIR ============
def generate_share_links(text: str, url: Optional[str] = None) -> ShareLinksResponse:
    encoded_text = urllib.parse.quote(text)
    encoded_url = urllib.parse.quote(url) if url else ""
    return ShareLinksResponse(
        whatsapp=f"https://wa.me/?text={encoded_text}",
        wechat=f"weixin://dl/chat?text={encoded_text}",
        facebook=f"https://www.facebook.com/sharer/sharer.php?quote={encoded_text}&u={encoded_url}" if url else f"https://www.facebook.com/sharer/sharer.php?quote={encoded_text}",
        tiktok=f"snssdk112://share?text={encoded_text}",
        instagram=f"instagram://library?AssetPath={encoded_text}",
        twitter=f"https://twitter.com/intent/tweet?text={encoded_text}&url={encoded_url}" if url else f"https://twitter.com/intent/tweet?text={encoded_text}",
        telegram=f"https://t.me/share/url?url={encoded_url}&text={encoded_text}" if url else f"https://t.me/share/url?text={encoded_text}",
        linkedin=f"https://www.linkedin.com/sharing/share-offsite/?url={encoded_url}" if url else f"https://www.linkedin.com/sharing/share-offsite/?url={encoded_text}",
        reddit=f"https://www.reddit.com/submit?title={encoded_text}&url={encoded_url}" if url else f"https://www.reddit.com/submit?title={encoded_text}",
        pinterest=f"https://pinterest.com/pin/create/button/?description={encoded_text}&url={encoded_url}" if url else f"https://pinterest.com/pin/create/button/?description={encoded_text}",
        snapchat=f"snapchat://share?text={encoded_text}",
        discord=f"https://discord.com/channels/@me?message={encoded_text}",
        email=f"mailto:?subject=Compartido desde ApoloXia&body={encoded_text}",
        copy_text=text
    )

async def send_whatsapp_message(phone_number: str, text: str) -> Dict:
    if not WHATSAPP_PHONE_NUMBER_ID or not WHATSAPP_ACCESS_TOKEN:
        return {"error": "WhatsApp Business API no configurada."}
    url = f"https://graph.facebook.com/v18.0/{WHATSAPP_PHONE_NUMBER_ID}/messages"
    headers = {"Authorization": f"Bearer {WHATSAPP_ACCESS_TOKEN}", "Content-Type": "application/json"}
    payload = {"messaging_product": "whatsapp", "to": phone_number, "type": "text", "text": {"body": text}}
    async with httpx.AsyncClient(timeout=15.0) as client:
        resp = await client.post(url, headers=headers, json=payload)
        return resp.json()

# ============ BÚSQUEDA WEB ============
async def search_exa_ai(query: str, max_results: int = 3) -> List[Dict]:
    if not EXA_API_KEY: return []
    headers = {"Authorization": f"Bearer {EXA_API_KEY}", "Content-Type": "application/json"}
    payload = {"query": query, "numResults": max_results, "type": "auto", "contents": {"text": True, "snippet": True}}
    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            resp = await client.post("https://api.exa.ai/search", headers=headers, json=payload)
            if resp.status_code != 200: return []
            data = resp.json()
            results = []
            for r in data.get("results", []):
                content = r.get("text", "") or r.get("snippet", "")
                img_urls = re.findall(r'https?://[^\s]+\.(?:jpg|jpeg|png|gif|webp|svg)', content, re.IGNORECASE)
                results.append({"title": r.get("title", ""), "url": r.get("url", ""), "content": content[:800],
                                "score": r.get("score", 0.5), "image_urls": img_urls[:2]})
            return results
    except: return []

async def search_mediastack(query: str, max_results: int = 3) -> List[Dict]:
    if not MEDIASTACK_API_KEY: return []
    params = {"access_key": MEDIASTACK_API_KEY, "keywords": query, "countries": "us,pa,es,mx,ar,co,cl,pe",
              "languages": "es,en", "limit": max_results, "sort": "published_desc"}
    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            resp = await client.get("http://api.mediastack.com/v1/news", params=params)
            if resp.status_code != 200: return []
            data = resp.json()
            return [{"title": a.get("title", ""), "description": a.get("description", ""), "url": a.get("url", ""),
                     "source": a.get("source", ""), "publishedAt": a.get("published_at", ""),
                     "image_url": a.get("image", ""), "category": a.get("category", ""), "language": a.get("language", "")}
                    for a in data.get("data", [])]
    except: return []

async def search_tavily(query: str, max_results: int = 3) -> List[Dict]:
    if not TAVILY_API_KEY: return []
    headers = {"Authorization": f"Bearer {TAVILY_API_KEY}", "Content-Type": "application/json"}
    search_query = f"{query} noticias actualidad" if any(k in query.lower() for k in ["noticia", "actual", "hoy"]) else query
    payload = {"query": search_query, "max_results": max_results, "search_depth": "basic", "include_answer": True}
    try:
        async with httpx.AsyncClient(timeout=10.0) as client:
            resp = await client.post("https://api.tavily.com/search", headers=headers, json=payload)
            if resp.status_code != 200: return []
            data = resp.json()
            results = []
            for r in data.get("results", []):
                content = r.get("content", "")
                img_urls = re.findall(r'https?://[^\s]+\.(?:jpg|jpeg|png|gif|webp|svg)', content, re.IGNORECASE)
                results.append({"title": r.get("title", ""), "url": r.get("url", ""), "content": content,
                                "score": r.get("score", 0), "image_urls": img_urls[:2]})
            if data.get("answer"):
                results.insert(0, {"title": "Respuesta destacada", "url": "", "content": data["answer"], "score": 1.0, "image_urls": []})
            return results
    except: return []

def needs_web_search(message: str) -> bool:
    return any(ind in message.lower() for ind in ["actualidad", "hoy", "reciente", "último", "news", "today", "latest", "2026", "precio de", "cotización", "noticia"])

# ============ LLAMADA A OPENROUTER ============
async def call_openrouter_api(messages: List[Dict], model_id: str, temperature: float = 0.7, max_tokens: int = 1024) -> str:
    if not OPENROUTER_API_KEY:
        raise HTTPException(500, "OPENROUTER_API_KEY no configurada")
    headers = {"Authorization": f"Bearer {OPENROUTER_API_KEY}", "Content-Type": "application/json",
               "HTTP-Referer": "https://theshieldtechnology.com", "X-OpenRouter-Title": "ApoloXia"}
    safe = min(max_tokens, 12000)
    payload = {"model": model_id, "messages": messages, "temperature": temperature, "max_tokens": safe}
    async with httpx.AsyncClient(timeout=100.0) as client:
        resp = await client.post("https://openrouter.ai/api/v1/chat/completions", headers=headers, json=payload)
        if resp.status_code != 200:
            err = resp.json()
            raise HTTPException(resp.status_code, f"OpenRouter: {err.get('error', {}).get('message', 'error')}")
        return resp.json()["choices"][0]["message"]["content"]

async def call_llm_pool_fallback(messages: List[Dict], temperature: float = 0.7, max_tokens: int = 1024) -> Dict[str, str]:
    if llm_pool is None:
        raise HTTPException(500, "freellmpool no disponible para fallback")
    system_prompt = ""
    user_message = ""
    for msg in messages:
        if msg["role"] == "system": system_prompt = msg["content"]
        elif msg["role"] == "user": user_message = msg["content"]
    if not user_message: raise HTTPException(400, "No hay mensaje de usuario")
    loop = asyncio.get_event_loop()
    reply = await loop.run_in_executor(None, lambda: llm_pool.ask(
        user_message, system=system_prompt if system_prompt else None,
        max_tokens=max_tokens, temperature=temperature))
    return {"text": reply.text, "provider": getattr(reply, "provider_id", "freellmpool")}

# ============ CONSTRUCCIÓN DE MENSAJES ============
def build_messages(user_id: str, conversation_id: str, user_message: str, agent_type: str, tier: str,
                   web_search_results: Optional[List[Dict]] = None, file_content: Optional[str] = None) -> List[Dict]:
    config = TIER_CONFIGS[tier]
    history = memory.get_memory(user_id, conversation_id, tier)
    try: agent = AgentType(agent_type) if agent_type else AgentType.GENERAL
    except ValueError: agent = AgentType.GENERAL
    if agent not in config.available_agents: agent = AgentType.GENERAL

    system = IDENTITY_OVERRIDE + AGENT_PROMPTS.get(agent, AGENT_PROMPTS[AgentType.GENERAL])
    system += "\n\n📢 Responde SIEMPRE en el MISMO IDIOMA del usuario.\n"
    system += INSTRUCCION_PROFUNDIDAD

    profile = get_business_profile(user_id)
    sector = profile.get("sector") or detect_business_sector(user_message)
    if sector:
        system += f"\n\n🏢 **CONTEXTO EMPRESARIAL — Sector: {sector.upper()}**\n{SECTOR_CONTEXT.get(sector, '')}\nAdapta tu lenguaje y ejemplos a este sector.\n"

    if profile:
        system += "\n=== PERFIL DE LA EMPRESA ===\n"
        if profile.get("company_name"): system += f"🏷️ Empresa: {profile['company_name']}\n"
        if profile.get("country"): system += f"🌎 País: {profile['country']}\n"
        if profile.get("tone"): system += f"🗣️ Tono: {profile['tone']}\n"

    if agent in (AgentType.DOCUMENTOS_EMPRESARIALES, AgentType.CONSULTOR_PYME, AgentType.PROCESOS_SOP,
                 AgentType.CUMPLIMIENTO_LEGAL_PA, AgentType.RRHH_RECLUTAMIENTO):
        system += "\n📄 Al final, añade: '(Puedes descargarlo en PDF o Word)'.\n"

    if tier in ["plus", "gt"]:
        system += "\n🔹 Responde con máximo detalle, múltiples secciones y análisis profundo.\n"

    if web_search_results:
        search_text = "\n\n=== INFORMACIÓN ACTUAL DE INTERNET ===\n"
        for i, res in enumerate(web_search_results[:4], 1):
            search_text += f"\n📌 {res.get('title', 'Sin título')}\n"
            if res.get('url'): search_text += f"🔗 {res['url']}\n"
            content = res.get('content', '') or res.get('description', '')
            if content: search_text += f"📝 {content[:400]}...\n"
            if res.get('source'): search_text += f"📰 {res['source']}\n"
        system += search_text

    if file_content:
        system += f"\n=== CONTENIDO DEL ARCHIVO ===\n{file_content[:2000]}\n"

    messages = [{"role": "system", "content": system}]
    for msg in history[-config.max_context_messages:]:
        messages.append({"role": msg["role"], "content": msg["content"]})
    messages.append({"role": "user", "content": user_message})
    return messages

# ============ FASTAPI APP ============
app = FastAPI(title="ApoloXia API", version="6.1.0-enterprise-openrouter")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

# ============ CHAT ============
@app.post("/chat", response_model=ChatResponse)
async def chat(request: ChatRequest):
    tier = request.tier.lower()
    if tier not in TIER_CONFIGS:
        raise HTTPException(400, f"Tier '{tier}' no válido")

    config = TIER_CONFIGS[tier]
    if not memory.check_daily_limit(request.user_id, tier):
        raise HTTPException(429, f"Límite diario alcanzado ({config.max_daily_responses})")

    conv_id = request.conversation_id or str(uuid.uuid4())
    use_web = request.use_web_search or needs_web_search(request.message)
    web_results = None

    if use_web and config.supports_web_search:
        tasks = []
        if TAVILY_API_KEY: tasks.append(search_tavily(request.message))
        if EXA_API_KEY: tasks.append(search_exa_ai(request.message))
        if MEDIASTACK_API_KEY: tasks.append(search_mediastack(request.message))
        if tasks:
            results = await asyncio.gather(*tasks, return_exceptions=True)
            combined = []
            for res in results:
                if isinstance(res, list) and res: combined.extend(res)
            web_results = combined[:8]

    messages = build_messages(request.user_id, conv_id, request.message,
                              request.agent_type or "general", tier, web_results, request.file_content)

    # === TOKENS ALTOS PARA ANÁLISIS PROFUNDO ===
    base_max_tokens = {"free": 2000, "plus": 6000, "gt": 10000}.get(tier, 2000)
    if request.agent_type in ["generador_sitios_web", "creador_paneles_ventas", "desarrollador_avanzado"]:
        max_tokens_for_agent = 12000
    else:
        max_tokens_for_agent = base_max_tokens
    safe_max_tokens = min(max_tokens_for_agent, 12000)
    temp = 0.7 if tier == "free" else 0.6

    # === SELECCIÓN DE MODELO OPENROUTER ===
    chosen = request.model_id or ""
    response_text = None
    model_used = None

    if chosen and chosen in OPENROUTER_MODELS:
        or_id = OPENROUTER_MODELS[chosen]
        try:
            print(f"🌐 OpenRouter directo: {or_id}")
            response_text = await call_openrouter_api(messages, or_id, temperature=temp, max_tokens=safe_max_tokens)
            model_used = or_id
        except Exception as e:
            print(f"⚠️ OpenRouter falló: {e}. Fallback a pool...")
            result = await call_llm_pool_fallback(messages, temperature=temp, max_tokens=safe_max_tokens)
            response_text = result["text"]; model_used = result["provider"]
    else:
        # Sin modelo elegido → usar el primero del tier
        default_key = config.available_models[0] if config.available_models else "or-gemma-4-31b"
        if default_key not in OPENROUTER_MODELS:
            default_key = "or-gemma-4-31b"
        or_id = OPENROUTER_MODELS[default_key]
        try:
            print(f"🌐 OpenRouter (default del tier): {or_id}")
            response_text = await call_openrouter_api(messages, or_id, temperature=temp, max_tokens=safe_max_tokens)
            model_used = or_id
        except Exception as e:
            print(f"⚠️ OpenRouter falló: {e}. Fallback a pool...")
            result = await call_llm_pool_fallback(messages, temperature=temp, max_tokens=safe_max_tokens)
            response_text = result["text"]; model_used = result["provider"]

    print(f"✅ Respuesta de: {model_used}")

    memory.add_message(request.user_id, conv_id, "user", request.message, request.agent_type)
    memory.add_message(request.user_id, conv_id, "assistant", response_text, request.agent_type)
    memory.increment_counter(request.user_id)

    detected_sector = get_business_profile(request.user_id).get("sector") or detect_business_sector(request.message)
    doc_agents = {"documentos_empresariales", "consultor_pyme", "procesos_sop", "cumplimiento_legal_pa", "rrhh_reclutamiento"}
    suggested_export = "pdf" if (request.agent_type in doc_agents or len(response_text) > 1200) else None

    return ChatResponse(
        response=response_text, agent_used=request.agent_type or "general", model_used=model_used,
        remaining_daily=memory.get_remaining(request.user_id, tier), web_search_used=bool(web_results),
        search_results=web_results, business_sector_detected=detected_sector, suggested_export=suggested_export,
    )

# ============ MODELOS DISPONIBLES ============
@app.get("/models")
async def list_models():
    or_list = [{"id": k, "provider": "openrouter", "api_id": v, "name": v.replace(":free", " (gratis)"),
                "free": True} for k, v in OPENROUTER_MODELS.items()]
    return {"openrouter": or_list, "total": len(or_list)}

@app.get("/user-config/{user_id}")
async def get_user_config(user_id: str):
    tier = memory.get_user_tier(user_id)
    config = TIER_CONFIGS[tier]
    user_cfg = memory.get_user_config(user_id)
    or_models = [{"id": k, "name": v.replace(":free", " (gratis)"), "provider": "openrouter", "free": True}
                 for k, v in OPENROUTER_MODELS.items()]
    return {"user_id": user_id, "tier": tier, "tier_name": config.name,
            "theme": user_cfg.get("theme", "dark"), "language": user_cfg.get("language", "es"),
            "notifications": user_cfg.get("notifications", True),
            "remaining_daily": memory.get_remaining(user_id, tier),
            "max_daily_responses": config.max_daily_responses,
            "available_models": or_models,
            "available_agents": [a.value for a in config.available_agents]}

@app.post("/user-config/{user_id}")
async def update_user_config(user_id: str, config: UserConfigUpdate):
    current = memory.get_user_config(user_id)
    if config.tier and config.tier in TIER_CONFIGS: memory.set_user_tier(user_id, config.tier)
    if config.theme: current["theme"] = config.theme
    if config.language: current["language"] = config.language
    if config.notifications is not None: current["notifications"] = config.notifications
    memory.set_user_config(user_id, current)
    return {"message": "Configuración actualizada", "user_id": user_id, "config": current}

@app.post("/upgrade-tier/{user_id}")
async def upgrade_tier(user_id: str, new_tier: str):
    if new_tier not in TIER_CONFIGS: raise HTTPException(400, "Tier no válido")
    memory.set_user_tier(user_id, new_tier)
    return {"message": f"Usuario {user_id} → {new_tier}", "tier": new_tier}

# ============ HEALTH ============
@app.get("/health")
async def health_check():
    return {"status": "ok", "version": "6.1.0-enterprise-openrouter",
            "openrouter_api": "configured" if OPENROUTER_API_KEY else "not set",
            "freellmpool": "active (fallback)" if llm_pool else "not available",
            "openrouter_models_count": len(OPENROUTER_MODELS),
            "identity": "ApoloXia · The Shield Technology · Panamá · Amelio Delgado"}

# ============ COMPARTIR ============
@app.post("/share/links", response_model=ShareLinksResponse)
async def get_share_links(share_req: ShareRequest):
    return generate_share_links(share_req.text, share_req.url)

@app.post("/share/send")
async def share_to_platform(share_req: ShareRequest):
    if share_req.platform.lower() == "whatsapp":
        if not share_req.user_phone_number: raise HTTPException(400, "Se requiere número")
        result = await send_whatsapp_message(share_req.user_phone_number, share_req.text)
        return {"platform": "whatsapp", "message": "Enviado", "result": result}
    links = generate_share_links(share_req.text, share_req.url)
    return {"platform": share_req.platform, "links": links.dict()}

@app.get("/share/links")
async def get_share_links_get(text: str, url: Optional[str] = None):
    return generate_share_links(text, url).dict()

# ============ CONVERSACIONES ============
@app.get("/conversations/{user_id}")
async def get_conversations(user_id: str):
    user_mem = memory.memories.get(user_id, {})
    convs = [{"conversation_id": cid, "message_count": len(msgs),
              "last_access": memory.last_access.get(user_id, {}).get(cid).isoformat() if memory.last_access.get(user_id, {}).get(cid) else None}
             for cid, msgs in user_mem.items()]
    return {"user_id": user_id, "conversations": convs}

@app.delete("/conversations/{user_id}/{conversation_id}")
async def delete_conversation(user_id: str, conversation_id: str):
    if user_id in memory.memories and conversation_id in memory.memories[user_id]:
        del memory.memories[user_id][conversation_id]
        if conversation_id in memory.last_access.get(user_id, {}): del memory.last_access[user_id][conversation_id]
        return {"message": "Conversación eliminada"}
    raise HTTPException(404, "No encontrada")

# ============ PERFIL DE EMPRESA ============
@app.get("/business-sectors")
async def list_business_sectors():
    return {"sectors": [{"id": k, "keywords": v, "context": SECTOR_CONTEXT.get(k, "")} for k, v in BUSINESS_SECTORS.items()],
            "total": len(BUSINESS_SECTORS)}

@app.get("/business-profile/{user_id}")
async def get_business_profile_endpoint(user_id: str):
    return {"user_id": user_id, "profile": get_business_profile(user_id)}

@app.post("/business-profile/{user_id}")
async def update_business_profile(user_id: str, upd: BusinessProfileUpdate):
    prof = get_business_profile(user_id)
    for field, value in upd.dict(exclude_none=True).items(): prof[field] = value
    if prof.get("sector") and prof["sector"] not in BUSINESS_SECTORS:
        raise HTTPException(400, f"Sector '{prof['sector']}' no soportado")
    set_business_profile(user_id, prof)
    return {"message": "Perfil actualizado", "user_id": user_id, "profile": prof}

# ============ EXPORTACIÓN PDF / WORD ============
def _markdown_to_flowables(content: str, styles, doc_width: float):
    flow = []
    lines = content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip(): flow.append(Spacer(1, 0.25*cm)); i += 1; continue
        if line.strip().startswith("```"):
            code_lines = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i]); i += 1
            flow.append(Preformatted("\n".join(code_lines) or "(vacío)", styles["Code"])); flow.append(Spacer(1, 0.3*cm)); i += 1; continue
        if line.startswith("### "): flow.append(Paragraph(line[4:], styles["H3"]))
        elif line.startswith("## "): flow.append(Paragraph(line[3:], styles["H2"]))
        elif line.startswith("# "): flow.append(Paragraph(line[2:], styles["H1"]))
        elif line.strip().startswith(("- ", "* ", "• ")):
            items = []
            while i < len(lines) and lines[i].strip().startswith(("- ", "* ", "• ")):
                items.append(ListItem(Paragraph(lines[i].strip()[2:], styles["Body"]))); i += 1
            flow.append(ListFlowable(items, bulletType="bullet")); continue
        elif re.match(r"^\d+\.\s", line.strip()):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i].strip()):
                items.append(ListItem(Paragraph(re.sub(r"^\d+\.\s", "", lines[i].strip()), styles["Body"]))); i += 1
            flow.append(ListFlowable(items, bulletType="1")); continue
        elif line.strip().startswith("|") and "|" in line.strip()[1:]:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c): rows.append(cells)
                i += 1
            if rows:
                t = Table(rows, colWidths=[doc_width/len(rows[0])]*len(rows[0]))
                t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1f2937")),
                    ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
                    ("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),8)]))
                flow.append(t); flow.append(Spacer(1, 0.3*cm))
            continue
        else:
            clean = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", line)
            clean = re.sub(r"\*(.+?)\*", r"<i>\1</i>", clean)
            flow.append(Paragraph(clean, styles["Body"]))
        i += 1
    return flow

@app.post("/export/pdf")
async def export_pdf(req: ExportRequest):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm,
                            title=req.title, author="ApoloXia")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="H1", parent=styles["Heading1"], fontSize=18, textColor=colors.HexColor("#0f172a"), spaceAfter=10))
    styles.add(ParagraphStyle(name="H2", parent=styles["Heading2"], fontSize=14, textColor=colors.HexColor("#1e293b"), spaceAfter=6))
    styles.add(ParagraphStyle(name="H3", parent=styles["Heading3"], fontSize=12, textColor=colors.HexColor("#334155"), spaceAfter=4))
    styles.add(ParagraphStyle(name="Body", parent=styles["BodyText"], fontSize=10.5, leading=15, spaceAfter=4))
    styles.add(ParagraphStyle(name="Code", parent=styles["Code"], fontSize=8.5, backColor=colors.HexColor("#f1f5f9"), borderPadding=6))
    flow = [Paragraph(req.title, styles["H1"]),
            Paragraph(f"<i>Generado por ApoloXia · {datetime.now():%Y-%m-%d %H:%M}</i>", styles["Body"]),
            Spacer(1, 0.5*cm)]
    if req.company_name: flow.append(Paragraph(f"<b>Empresa:</b> {req.company_name}", styles["Body"])); flow.append(Spacer(1, 0.3*cm))
    flow.extend(_markdown_to_flowables(req.content, styles, A4[0]-4*cm))
    doc.build(flow); buf.seek(0)
    filename = re.sub(r"[^\w\-]+", "_", req.title)[:60] or "documento_apoloxia"
    return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}.pdf"'})

@app.post("/export/docx")
async def export_docx(req: ExportRequest):
    doc = Document()
    style = doc.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(11)
    h = doc.add_heading(req.title, level=0); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Generado por ApoloXia · {datetime.now():%Y-%m-%d %H:%M}")
    run.italic = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    if req.company_name: doc.add_paragraph(f"Empresa: {req.company_name}").runs[0].bold = True
    doc.add_paragraph()
    lines = req.content.split("\n"); i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.strip().startswith("```"):
            code_lines = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i]); i += 1
            p = doc.add_paragraph(); r = p.add_run("\n".join(code_lines) or "(vacío)"); r.font.name = "Consolas"; r.font.size = Pt(9); i += 1; continue
        if line.startswith("### "): doc.add_heading(line[4:], level=3)
        elif line.startswith("## "): doc.add_heading(line[3:], level=2)
        elif line.startswith("# "): doc.add_heading(line[2:], level=1)
        elif line.strip().startswith(("- ", "* ", "• ")): doc.add_paragraph(line.strip()[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", line.strip()): doc.add_paragraph(re.sub(r"^\d+\.\s", "", line.strip()), style="List Number")
        elif line.strip().startswith("|") and "|" in line.strip()[1:]:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c): rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row): t.cell(ri, ci).text = val
            continue
        elif line.strip():
            p = doc.add_paragraph()
            for chunk in re.split(r"(\*\*.+?\*\*)", line):
                if chunk.startswith("**") and chunk.endswith("**"):
                    run = p.add_run(chunk[2:-2]); run.bold = True
                else: p.add_run(chunk)
        else: doc.add_paragraph("")
        i += 1
    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    filename = re.sub(r"[^\w\-]+", "_", req.title)[:60] or "documento_apoloxia"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": f'attachment; filename="{filename}.docx"'})

# ============ ARCHIVOS ESTÁTICOS ============
def _get_file_path(filename: str) -> str:
    if os.path.exists(filename) and os.path.isfile(filename): return filename
    parent = os.path.join(os.path.dirname(__file__), "..", filename)
    if os.path.exists(parent) and os.path.isfile(parent): return parent
    return filename

@app.get("/")
async def serve_index(): return FileResponse(_get_file_path("index.html"))

@app.get("/chat")
async def serve_chat(): return FileResponse(_get_file_path("chat.html"))

@app.get("/apoloxia.code.html")
async def serve_code(): return FileResponse(_get_file_path("apoloxia.code.html"))

@app.get("/{filename}")
async def serve_static_file(filename: str):
    api_routes = {"chat","models","agents","health","conversations","tier-info","user-config","upgrade-tier",
                  "search-web","share","export","business-sectors","business-profile"}
    if filename in api_routes or filename.startswith("apoloxia.code.html"): raise HTTPException(404, "Not found")
    if filename.startswith(".") or ".." in filename: raise HTTPException(403, "Forbidden")
    fp = _get_file_path(filename)
    if os.path.exists(fp) and os.path.isfile(fp): return FileResponse(fp)
    raise HTTPException(404, "Archivo no encontrado")

# ============ MAIN ============
if __name__ == "__main__":
    print("🚀 ApoloXia Server v6.1.0 ENTERPRISE (Solo OpenRouter)")
    print(f"🌐 OpenRouter gratuitos: {len(OPENROUTER_MODELS)}")
    print(f"🧠 freellmpool fallback: {'✅' if llm_pool else '❌'}")
    print("=" * 60)
    uvicorn.run(app, host="0.0.0.0", port=8000)